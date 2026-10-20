#!/usr/bin/env python3
"""Convert archive DOCX to docs/*.md."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT / "archive" / "canvas-submissions"
DOCS = ROOT / "docs"

SLUGS: dict[str, str] = {
    "Introduction to the PlayStation Network Breach.docx": "introduction-playstation-network-breach.md",
    "Introduction to the PlayStation Network Breach-1.docx": "introduction-playstation-network-breach-canvas-copy.md",
    "Milestone two .docx": "milestone-two.md",
    "Milestone Three.docx": "milestone-three.md",
    "Final Assignment.docx": "final-assignment.md",
    "Final Assignment-1.docx": "final-assignment-canvas-copy.md",
}


def table_to_md(table) -> list[str]:
    rows = [[c.text.strip().replace("\n", " ") for c in row.cells] for row in table.rows]
    if not rows:
        return []
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join("---" for _ in rows[0]) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return lines


def convert(src: Path, dest: Path, title: str) -> None:
    doc = Document(src)
    rel = src.relative_to(ROOT).as_posix().replace(" ", "%20")
    lines: list[str] = [
        f"# {title}",
        "",
        f"**Source:** [`{src.name}`](../{rel})  ",
        "**Course:** Cyber Crime Investigation (EMU)  ",
        "**Author:** Julio Barros",
        "",
    ]
    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            lines.append("")
            continue
        style = (para.style.name or "").lower()
        if "heading 1" in style:
            lines.extend(["", f"## {t}", ""])
        elif "heading 2" in style:
            lines.extend(["", f"### {t}", ""])
        elif "heading 3" in style:
            lines.extend(["", f"#### {t}", ""])
        elif "list" in style:
            lines.append(f"- {t}")
        else:
            lines.append(t)
            lines.append("")
    for table in doc.tables:
        lines.extend(["", *table_to_md(table), ""])
    body = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"
    dest.write_text(body, encoding="utf-8")


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    for name, slug in SLUGS.items():
        src = ARCHIVE / name
        if not src.exists():
            print(f"skip {name}")
            continue
        convert(src, DOCS / slug, name.removesuffix(".docx"))
        print(f"wrote docs/{slug}")


if __name__ == "__main__":
    main()
